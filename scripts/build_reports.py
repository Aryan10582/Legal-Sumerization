"""Compose the supplied MIT cover and actual IEEE ProComm styles into editable DOCX."""
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
from collections import Counter
import json,re,hashlib
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path.cwd(); COVER=Path('templates/ML_Front_Page.docx'); IEEE=Path('templates/Full-Paper-template.docx')
TITLE='AI-Based Summarization and Clause Classification of Indian Legal Documents Using NLP and Transformer Models'
BODYFONT='Times New Roman'

def set_text(p,text,size=None,bold=None,italic=False,font=BODYFONT):
    p.clear(); r=p.add_run(text); r.font.name=font
    if size: r.font.size=Pt(size)
    if bold is not None:r.bold=bold
    r.italic=italic
    return p

def cover(kind):
    d=Document(COVER)
    # Materialize cover typography before body style replacement.
    for p in d.paragraphs:
        for r in p.runs:
            if r.text.strip() and not r.font.name: r.font.name='Garamond'
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs: r.font.name='Garamond';r.font.size=Pt(10)
    p=d.paragraphs
    set_text(p[2],'Mini Project '+('Report' if kind=='final' else 'Problem Statement'),22,True,font='Garamond')
    set_text(p[4],'Machine Learning (CSE 3125)',20,True,font='Garamond')
    set_text(p[8],TITLE,17,True,font='Garamond')
    set_text(p[16],'Guide(s): [CONFIRM NAME(S), SPELLING AND DESIGNATION]',11,True,font='Garamond')
    set_text(p[17],'Synopsis: Dr. Rani Ooman Panicker; Dr. H Satyajeet Sharma',9,font='Garamond')
    set_text(p[18],'Front-page source: Dr. Rani Oomman Panicker',9,font='Garamond')
    set_text(p[23],'October 2026  |  Academic Year 2026–27',16,True,font='Garamond')
    # Existing five-column student table and original logo/border stay in place.
    t=d.tables[0]
    while len(t.rows)>4:t._tbl.remove(t.rows[-1]._tr)
    content=[['Vedant jitani','240962492','41','V, CSE (AI & ML)','C'],
             ['[STUDENT 2 NAME]','[REGISTRATION 2]','[ROLL 2]','[SEMESTER & BRANCH 2]','[SECTION 2]'],
             ['[STUDENT 3 NAME]','[REGISTRATION 3]','[ROLL 3]','[SEMESTER & BRANCH 3]','[SECTION 3]']]
    for row,values in zip(t.rows[1:],content):
        for cell,val in zip(row.cells,values):
            cell.text=val
            for para in cell.paragraphs:
                para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:run.font.name='Garamond';run.font.size=Pt(8.5)
    for row in t.rows:
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for i in [5,6,7,9,10,11,14,15,19,20]:
        p[i].paragraph_format.space_after=Pt(2)
        for r in p[i].runs:r.font.size=Pt(8)
    d.core_properties.title=TITLE
    d.core_properties.subject='Machine Learning mini-project; pilot study'
    d.core_properties.author='Three-student team (two names pending)'
    return d

def adopt_ieee_styles(d):
    template=Document(IEEE)
    target=d.styles.element
    for s in template.styles.element:
        sid=s.get(qn('w:styleId'))
        if not sid:continue
        old=next((x for x in target if x.tag==qn('w:style') and x.get(qn('w:styleId'))==sid),None)
        if old is not None:target.replace(old,deepcopy(s))
        else:target.append(deepcopy(s))
    # Body normal from the actual template; cover content carries explicit direct format.
    return d

def set_section(sec,columns=1):
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(1);sec.bottom_margin=Inches(1)
    sec.left_margin=Inches(1210/1440);sec.right_margin=Inches(1210/1440)
    cols=sec._sectPr.xpath('./w:cols')
    if not cols:
        c=OxmlElement('w:cols');sec._sectPr.append(c)
    else:c=cols[0]
    c.set(qn('w:num'),str(columns));c.set(qn('w:space'),'461' if columns==2 else '720')
    for border in sec._sectPr.xpath('./w:pgBorders'):
        sec._sectPr.remove(border)
    # Template does not paginate body.
    sec.different_first_page_header_footer=False

def heading(d,title):
    p=d.add_paragraph(style='Heading 1');p.add_run(title);p.paragraph_format.keep_with_next=True;return p
def para(d,text,style='Body Text Indent'):
    p=d.add_paragraph(style=style);p.add_run(text);p.paragraph_format.widow_control=True;return p
def table(d,headers,rows,widths=None):
    t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER
    # Supplied cover/template have no generic Table Grid style; keep real table geometry.
    for j,h in enumerate(headers):
        c=t.rows[0].cells[j];c.text=h
        for p in c.paragraphs:
            for r in p.runs:r.bold=True;r.font.name=BODYFONT;r.font.size=Pt(8)
    for values in rows:
        row=t.add_row()
        for c,val in zip(row.cells,values):
            c.text=str(val)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(0)
                for r in p.runs:r.font.name=BODYFONT;r.font.size=Pt(7.6)
    if widths:
        for row in t.rows:
            for cell,w in zip(row.cells,widths):cell.width=Inches(w)
    for row in t.rows:
        for c in row.cells:c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return t

abstract=(
    'Long court judgments and commercial agreements are difficult to scan quickly, yet automated outputs can omit controlling facts or misidentify clauses. This project implements two separate tasks in one local interface: summarization of Indian Supreme Court judgments and multi-label classification of English commercial contract passages. For summarization, IN-Abs supplies judgment–headnote pairs. A full-document TF-IDF and maximum-marginal-relevance selector provides an extractive baseline and supplies bounded context to a fine-tuned T5-small generator. For classification, eight CUAD question-answer span categories are converted to passage labels; a TF-IDF logistic-regression baseline is compared with a fine-tuned LEGAL-BERT-small classifier. Documents are separated before segmentation, duplicate groups are audited, and validation selects model settings without test data. In the executed CPU pilot, LEGAL-BERT reached 0.447 macro F1 versus 0.385 for logistic regression on 2,404 held-out passages. The extractive summary baseline achieved ROUGE-1/2/L F1 of 0.216/0.120/0.145, above T5 at 0.149/0.080/0.105 on twenty judgments. These limited results support a working demonstration, not Indian-contract accuracy or legally faithful automated advice.'
)
intro_statement=(
    'Indian Supreme Court judgments can be lengthy, and a reader must still identify the issues, reasoning and final order. Commercial contracts present a different reading problem: important obligations may be distributed across many passages and a passage can express several clause types. This mini-project builds a local tool that places its output beside the source text so a reader can verify it. The first task summarizes Indian judgments using the public IN-Abs collection of judgment–headnote pairs. The second locates eight types of clauses using CUAD, a public expert-annotated collection of English commercial contracts from United States filings. CUAD is therefore a method benchmark, not evidence that the classifier has been validated on Indian contracts. We compare an extractive summary with a fine-tuned small generator, and a TF-IDF clause classifier with a fine-tuned legal-language encoder. Original documents are split before passages are made; duplicate and annotation-offset checks limit avoidable leakage or label errors. Evaluation uses held-out ROUGE and multi-label precision, recall and F1, with visible examples of omissions and false detections. The proposed use is research assistance under human review, not a legal conclusion or replacement for reading the judgment or agreement.'
)

def paper_header(d):
    s=d.add_section(WD_SECTION_START.NEW_PAGE);set_section(s,1)
    for part in [s.header,s.even_page_header,s.first_page_header,s.footer,s.even_page_footer,s.first_page_footer]:
        part.is_linked_to_previous=False
        for child in list(part._element):part._element.remove(child)
        part._element.append(OxmlElement('w:p'))
    # The cover's line border is a header drawing in the supplied template. Word
    # drops that drawing when later sections use independent blank headers, so
    # retain the same cover framing with an explicit page border on section 0.
    border=OxmlElement('w:pgBorders');border.set(qn('w:offsetFrom'),'page')
    for side in ['top','left','bottom','right']:
        edge=OxmlElement('w:'+side)
        edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'8')
        edge.set(qn('w:space'),'24');edge.set(qn('w:color'),'000000')
        border.append(edge)
    d.sections[0]._sectPr.append(border)
    p=d.add_paragraph(style='paper title');p.add_run(TITLE)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(10)
    t=table(d,['Vedant jitani','[STUDENT 2 NAME]','[STUDENT 3 NAME]'],[
      ['MIT Manipal\nReg. 240962492','MIT Manipal\n[REGISTRATION 2]','MIT Manipal\n[REGISTRATION 3]']])
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=d.add_paragraph(style='Abstract');p.add_run('Abstract - '+abstract)
    p=d.add_paragraph(style='Abstract');p.add_run('Index Terms - clause classification, Indian judgments, legal natural language processing, multi-label learning, summarization')
    s=d.add_section(WD_SECTION_START.CONTINUOUS);set_section(s,2)

def make_final():
    d=adopt_ieee_styles(cover('final'));paper_header(d)
    heading(d,'Introduction')
    para(d,'A long judgment may mix procedural history, arguments, reasons and the order. A contract can express several obligations in one passage. A usable tool should shorten reading while preserving a route back to the source. This project joins two related reading aids in one interface but trains and evaluates them independently. Only the summarization task uses Indian documents; CUAD clause labels come from non-Indian commercial contracts. A judgment repository by itself does not supply supervised headnotes, so paired IN-Abs was selected. Automatic outputs require human verification, especially for legal effect and dispositions [1], [2].')
    para(d,'The research question is whether small, reproducible baselines and genuinely fine-tuned transformers can be compared fairly on each task under CPU limits. The contribution is a measured pilot with document-level split control, transparent weak-label conversion, visible source coverage and error examples. It does not claim a new legal method or an Indian contract benchmark.')
    heading(d,'Literature Survey')
    para(d,'Legal language adaptation and rhetorical-role work established the value of domain context and structure [3], [4]. Greek and Chinese case-law studies show that language and judgment format affect transfer [5], [6]. Indian legal extraction has used neural sentence relevance [7]. Entailment reranking can improve generated summaries, yet entailment scores do not certify legal correctness [8]. Deep clustering and extractive ensembles improve content selection on their own benchmarks [9], [10].')
    para(d,'Generative-model comparisons still report omissions or unsupported content [11]. Structural hybrid approaches combine extraction and abstraction [12]–[14], motivating our explicit selected-context trace. Contract-risk BERT systems [15], multi-label augmentation [16] and abusive-clause assistance [17] use distinct taxonomies; their scores cannot be imported to CUAD. Indian judgment-domain classification [18], legal attribute extraction [19] and Thai multi-intent classification [20] further show that task and jurisdiction determine label meaning. Our gap is a locally reproducible dual-task teaching system with candid scope and data-conversion evidence, not a priority claim.')
    para(d,'Table I compresses the comparison; the twenty-study analytical record in LITERATURE_REVIEW.md states each data setting, method, reported finding, limitation and project implication. Older [1], [3], [4] are retained as foundations; the other seventeen are 2020 onward.')
    table(d,['Study group','Finding and design consequence'],[
        ['Legal/Indian summaries [1]–[4], [7]','Domain wording and roles matter; retain source and test Indian paired data.'],
        ['Selection/generation [5], [6], [8]–[14]','Selection can omit holding; compare extractive and generated outputs.'],
        ['Contract/label models [15]–[20]','Taxonomies differ; test eight explicit labels with per-class metrics.']],[1.05,2.05])
    para(d,'TABLE I. Literature groups and design implications. Studies are cited by final journal-reference number.','Caption')
    heading(d,'Objectives')
    para(d,'1) Build a source-linked Indian judgment summarizer and compare an extractive method with a fine-tuned generator. 2) Convert verified CUAD spans to eight multi-label passage categories and compare TF-IDF logistic regression with a fine-tuned legal encoder. 3) evaluate both tasks on isolated held-out documents, inspect errors, and expose their limits in an exportable local demonstration.')
    heading(d,'Data Source')
    para(d,'IN-Abs, released by Shukla et al. on Zenodo under CC BY 4.0, contains English Indian Supreme Court judgments and abstractive headnotes. Its published 7,130 pairs include 7,030 train and 100 official test. Two training files are empty, leaving 7,128 usable pairs. CUAD v1, from The Atticus Project under CC BY 4.0, contains 510 English commercial contracts from SEC filings, 41 expert question-answer annotation types and an official 408/102 train/test partition. Its contracts are non-Indian. Source URLs, hashes and non-journal provenance are in SOURCE_RECORD.md and DATA_CARD.md. We did not scrape Indian Kanoon or infer that NJDG supplied supervised pairs.')
    table(d,['Data / split','Train','Validation','Test'],[
        ['IN-Abs pool','5,949','1,050','100'],['IN-Abs pilot','64','12','20'],
        ['CUAD pool contracts','345','61','102'],['CUAD pilot contracts','40','10','10'],
        ['CUAD pilot segments','1,636','2,639','2,404']],[1.23,.59,.68,.6])
    para(d,'TABLE II. Retained pool after duplicate grouping versus executed pilot. The CUAD training passage count reflects negative downsampling; test retains all passages.','Caption')
    para(d,'The eight selected labels are Governing Law, Non-Compete, Exclusivity, Anti-Assignment, Termination For Convenience, Cap On Liability, Audit Rights and Post-Termination Services. Multiple categories may overlap one passage. The task identifies category presence, not validity, enforceability or risk. A passage without one of these annotations is not proven neutral.')
    heading(d,'Methodology')
    para(d,'Raw ZIPs are checksum-tested. IN-Abs source/summary pairs and CUAD untouched contexts/answer offsets are validated. All CUAD answer offsets matched their text. Documents, not passages, are grouped by normalized exact hash and an approximate five-word-shingle sketch, then split. Official test takes precedence. The sketch found one exact and one near-duplicate CUAD link, and 27 exact/two near IN-Abs links. Near matching is approximate. Exact repeated cross-split CUAD passages were further removed from lower-priority splits. Validation comes from official training groups; seed 42 fixes sampling. Fig. 1 shows both independent branches.')
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture('reports/figures/architecture.png',width=Inches(3.0))
    para(d,'Fig. 1. Local architecture. The judgment and contract branches share document input and source-linked display, but have independent supervision, models and metrics. Original figure.','Caption')
    para(d,'CUAD is formatted as question answering. We reconstruct each contract and mark a token-bounded passage positive for a category when it overlaps a verified answer span. A passage can have several 1s in an eight-element vector. Line-aware splitting preserves character offsets and meaningful negation, dates, names and amounts; recursive splitting enforces 256 LEGAL-BERT tokens. Training negatives are capped at three per positive; validation and test keep their natural imbalance. No softmax or mutually exclusive label assumption is made.')
    para(d,'For judgments, a full-text TF-IDF centroid ranks sentence units. Maximum marginal relevance penalizes redundancy, and selected units return to source order. The extractive baseline emits selected text within 128 T5 tokens. The same selector supplies at most 512 T5 source tokens; the model generates at most 128 new tokens with two beams. Selected spans and character coverage are saved. The generator cannot see omitted text, so this approach bounds compute rather than solving long-document completeness.')
    heading(d,'Implementation Details')
    para(d,'Python 3.12.14, PyTorch 2.14.1+cpu, Transformers 5.19.0, scikit-learn 1.9.1 and Streamlit 1.65.0 were installed together; pip check passed. Hardware was AMD64 Windows 11, 15.3 GiB RAM and four CPU threads, without CUDA. The classifier baseline uses unigram/bigram TF-IDF (40,000-feature cap) and eight balanced one-vs-rest logistic regressions. LEGAL-BERT-small is a 35.07-million-parameter encoder with an eight-logit head, trained end-to-end with weighted binary cross-entropy, AdamW, learning rate 5×10⁻⁵, batch eight, three epochs and 256-token passages. The best validation macro-F1 epoch was 3; one global threshold of 0.90 came from a 0.10–0.90 validation grid. TF-IDF threshold was 0.75.')
    para(d,'T5-small is a 60.51-million-parameter encoder-decoder fine-tuned for one epoch on 64 pairs, batch two, gradient accumulation two, learning rate 10⁻⁴ and teacher-forced validation-loss selection. All 64 train and 12 validation reference headnotes exceeded 128 target tokens and were explicitly shortened for training; test references remained complete. Mean training source-character coverage after selection was 16.0%. LEGAL-BERT training took 723.75 s; T5 took 152.02 s including model/data setup. Checkpoints store weights, tokenizer, settings, selected epoch and dataset hashes. The CLI does not save optimizer state for exact interrupted resume.')
    para(d,'The local GUI accepts pasted text or UTF-8 TXT, DOCX and searchable PDF up to 20 MB, displays each result beside the source, and exports TXT, CSV or JSON. Scanned PDFs need external OCR; extraction reports missing text. User text is processed locally. The trained classifier never runs a random head; a missing checkpoint produces an error.')
    heading(d,'Results and Discussion')
    para(d,'On the same 2,404 held-out CUAD passages from ten contracts, macro precision/recall/F1 for TF-IDF were 0.419/0.386/0.385 and for LEGAL-BERT 0.435/0.565/0.447. Weighted F1 was 0.409 versus 0.509; exact subset accuracy 0.970 versus 0.972. Hamming losses were 0.00374 and 0.00364. The test has 2,340 passages without any selected label. An always-zero classifier would reach 0.973 subset accuracy but zero macro F1, so the high accuracy cannot establish good clause detection.')
    table(d,['Model','Macro F1','Weighted F1','Subset acc.'],[
        ['TF-IDF + LR','0.385','0.409','0.970'],['LEGAL-BERT','0.447','0.509','0.972']],[1.26,.61,.72,.61])
    para(d,'TABLE III. Actual multi-label held-out CUAD pilot scores (0–1). Thresholds selected on validation only.','Caption')
    para(d,'Per-class LEGAL-BERT F1 ranged from 0.818 on Governing Law (ten positives) and 0.800 on Anti-Assignment (eleven) to 0 on Post-Termination Services (eight) and 0.182 on Exclusivity (three). It made 20 false positives for Cap On Liability and missed five of six Termination For Convenience passages. One Cap false positive was a standalone sentence that limitations “shall survive”; it did not state a numerical cap. One Post-Termination miss described compensation continuing for five years after termination. These illustrate partial-context and minority-label problems. Scores are uncalibrated and the ten-contract sample gives little basis for generalization.')
    para(d,'For twenty held-out Indian judgments, the extractive baseline achieved ROUGE-1/2/L F1 0.216/0.120/0.145, and fine-tuned T5 achieved 0.149/0.080/0.105 against full headnotes. The methods used the same documents and 128-token output budget. Inference on four CPU threads took 3.59 s total for extraction and 36.52 s for T5, excluding model loading. These are single measured runs, not service latency promises.')
    table(d,['Model','R-1 F1','R-2 F1','R-L F1'],[
        ['TF-IDF/MMR','0.216','0.120','0.145'],['Fine-tuned T5','0.149','0.080','0.105']],[1.26,.61,.61,.72])
    para(d,'TABLE IV. Actual document-mean ROUGE F1 on twenty IN-Abs test pairs; Porter stemming, rouge-score 0.1.2, full references.','Caption')
    para(d,'In held-out judgment 3542, the reference says the appeals were allowed and contempt proceedings dropped after faulting acceptance of an apology without a contempt finding. T5 summarized the letter and an argument but omitted the final disposition and key reason. The extractive output retained part of the apology reasoning but ended with an incomplete “The High.” fragment and still omitted the final order. This is a concrete content and segmentation failure. ROUGE rewards lexical overlap but does not establish faithfulness; no expert factuality score was measured. The selected input covered about 16.7% of this source, and all training references exceeded the target cap.')
    heading(d,'Conclusion and Future Scope')
    para(d,'The implemented local pilot performs two distinct reading tasks with tested data provenance, checkpoints and a usable interface. LEGAL-BERT improved macro F1 over a lexical baseline in this CUAD subset, while the extractive judgment baseline exceeded the one-epoch T5 generator on all three ROUGE scores. Neither result proves Indian contract performance or reliable legal summaries. Human source comparison is essential. Next experiments should run the full retained splits, obtain consented expert-labeled Indian contracts, test OCR and multilingual extraction, extend or structure the judgment input, evaluate legal issue/disposition faithfulness with expert review, and calibrate thresholds on broader validation data. These extensions require new measurements rather than retroactive claims about this pilot.')
    heading(d,'References')
    refs=Path('research/bibliography-ieee.txt').read_text(encoding='utf-8').strip().split('\n\n')
    assert len(refs)==20,len(refs)
    for ref in refs:
        p=d.add_paragraph();p.style='Normal';p.paragraph_format.first_line_indent=Inches(-.17);p.paragraph_format.left_indent=Inches(.17);p.paragraph_format.space_after=Pt(5)
        r=p.add_run(ref);r.font.name=BODYFONT;r.font.size=Pt(8.5)
    path=Path('reports/Final_Report.docx');d.save(path)
    return path

def make_initial():
    d=adopt_ieee_styles(cover('initial'))
    sec=d.add_section(WD_SECTION_START.NEW_PAGE);set_section(sec,1)
    p=d.add_paragraph(style='Normal');set_text(p,'Initial Mini-Project Problem Statement',16,True)
    d.add_paragraph(TITLE)
    heading(d,'Introduction and Real-World Problem')
    para(d,intro_statement,'Normal')
    heading(d,'Keywords')
    para(d,'clause classification; Indian judgments; legal NLP; multi-label learning; summarization.','Normal')
    heading(d,'Objectives')
    for item in [
        '1. Summarize Indian Supreme Court judgments from verified judgment–headnote pairs, comparing an extractive baseline and a fine-tuned generator.',
        '2. Convert eight verified CUAD clause spans into multi-label passages and compare a lexical classifier with a fine-tuned legal encoder.',
        '3. Report held-out task-specific scores and source-linked errors in a local, exportable interface.']:
        para(d,item,'Normal')
    heading(d,'Data Sources and Scope')
    para(d,'IN-Abs: Shukla et al., Zenodo record 7152317, DOI 10.5281/zenodo.7152317, CC BY 4.0; English Indian Supreme Court judgment–headnote pairs. CUAD v1: The Atticus Project, https://www.atticusprojectai.org/cuad, CC BY 4.0; English SEC commercial contracts with expert QA spans. CUAD is not an Indian contract benchmark. Operational provenance, checksums and label conversion are supplied in docs/DATA_CARD.md and docs/SOURCE_RECORD.md.','Normal')
    heading(d,'Three-Student Team and Guide')
    table(d,['Student','Registration','Roll','Semester/Branch','Section'],[
      ['Vedant jitani','240962492','41','V CSE (AI & ML)','C'],
      ['[STUDENT 2 NAME]','[REGISTRATION 2]','[ROLL 2]','[CONFIRM]','[CONFIRM]'],
      ['[STUDENT 3 NAME]','[REGISTRATION 3]','[ROLL 3]','[CONFIRM]','[CONFIRM]']])
    para(d,'Guide(s): [CONFIRM NAME(S), SPELLING AND DESIGNATION]. The supplied synopsis and front-page spell Dr. Rani Panicker’s middle name differently and the synopsis additionally lists Dr. H Satyajeet Sharma.','Normal')
    path=Path('reports/Initial_Problem_Statement.docx');d.save(path)
    return path

def inspect(files):
    result={'abstract_words':len(abstract.split()),'initial_introduction_words':len(intro_statement.split()),'outputs':{}}
    assert 150<=result['abstract_words']<=200,result
    assert 150<=result['initial_introduction_words']<=200,result
    with ZipFile(COVER) as z: src=[hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')]
    for path in files:
        d=Document(path)
        with ZipFile(path) as z: imgs=[hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')]
        assert set(src)<=set(imgs),'Institution logo missing'
        assert len(d.tables[0].rows)==4,'Cover not exactly three student rows'
        assert TITLE in '\n'.join(p.text for p in d.paragraphs)
        result['outputs'][path.name]={'bytes':path.stat().st_size,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
           'cover_rows':len(d.tables[0].rows),'images':len(imgs),'logo_preserved':True,
           'paragraphs':len(d.paragraphs),'sections':len(d.sections),'tables':len(d.tables),
           'body_columns':[x._sectPr.xpath('./w:cols')[0].get(qn('w:num')) for x in d.sections[1:]]}
    Path('reports/structural_QA.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
    print(json.dumps(result,indent=2))

if __name__=='__main__':
    files=[make_final(),make_initial()]
    inspect(files)
