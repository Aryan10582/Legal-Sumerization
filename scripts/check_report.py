from pathlib import Path
from hashlib import sha256
import json,re
from docx import Document
from pypdf import PdfReader

required=['Introduction','Literature Survey','Objectives','Data Source','Methodology','Implementation Details','Results and Discussion','Conclusion and Future Scope','References']
results={}
for stem,expected_pages in [('Final_Report',5),('Initial_Problem_Statement',2)]:
    docx=Path('reports',stem+'.docx');pdf=Path('reports',stem+'.pdf')
    d=Document(docx);p=PdfReader(pdf)
    assert len(p.pages)==expected_pages,(stem,len(p.pages))
    assert all(round(float(page.mediabox.width))==612 and round(float(page.mediabox.height))==792 for page in p.pages)
    assert len(d.tables[0].rows)==4
    text='\n'.join(x.text for x in d.paragraphs)
    if stem=='Final_Report':
        positions=[text.find(x) for x in required];assert positions==sorted(positions) and all(x>=0 for x in positions),positions
        pdftext='\n'.join(page.extract_text() for page in p.pages)
        assert all(f'[{n}]' in pdftext for n in range(1,21))
        assert len(p.pages)-1<=8
    else:
        assert 'Introduction and Real-World Problem' in text and 'Vedant jitani' in d.tables[0].rows[1].cells[0].text
    results[stem]={'docx_bytes':docx.stat().st_size,'pdf_bytes':pdf.stat().st_size,
                   'docx_sha256':sha256(docx.read_bytes()).hexdigest(),
                   'pdf_sha256':sha256(pdf.read_bytes()).hexdigest(),
                   'pages':len(p.pages),'cover_student_rows':3,
                   'page_size':'US Letter 612 x 792 pt'}
Path('reports/final_QA.json').write_text(json.dumps(results,indent=2),encoding='utf-8')
print(json.dumps(results,indent=2))
