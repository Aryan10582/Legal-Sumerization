"""Read supplied references and fetch small provenance records (no scraping)."""
from pathlib import Path
import json, hashlib, urllib.request, sys
from docx import Document
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
def fetch(url, dest):
    req = urllib.request.Request(url, headers={"User-Agent": "MIT-ML-mini-project/1.0"})
    with urllib.request.urlopen(req, timeout=90) as r:
        data = r.read()
    path = ROOT / dest
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    print(dest, len(data), hashlib.sha256(data).hexdigest())

def inspect():
    for name in ["mimi-project-Synopsis_pcap.pdf", "Mini-project-guidelines.pdf"]:
        p = Path("D:/3rd year/New folder") / name
        text = "\n\n".join(f"PAGE {i+1}\n{page.extract_text()}" for i,page in enumerate(PdfReader(p).pages))
        (ROOT / "research" / (p.stem + ".txt")).write_text(text, encoding="utf-8")
    for p in [Path("D:/3rd year/New folder/ML_Front_Page.docx"), ROOT / "tmp/reference/Full-Paper-template.docx"]:
        d = Document(p)
        out = [f"SHA256 {hashlib.sha256(p.read_bytes()).hexdigest()}"]
        out += [f"P{i} [{x.style.name}] {x.text!r}" for i,x in enumerate(d.paragraphs)]
        out += [f"TABLE {i}: " + str([[c.text for c in row.cells] for row in t.rows]) for i,t in enumerate(d.tables)]
        out += [s._sectPr.xml for s in d.sections]
        out += [f"STYLE {s.name}: {s.element.xml}" for s in d.styles if s.type == 1]
        from zipfile import ZipFile
        from lxml import etree
        z = ZipFile(p)
        out += ["ALL TEXT " + str(etree.fromstring(z.read('word/document.xml')).xpath('//*[local-name()="t"]/text()'))]
        (ROOT/"research"/(p.stem + "-inspection.txt")).write_text("\n".join(out),encoding="utf-8")
        print(p.name, len(d.paragraphs), "paragraphs", len(d.tables), "tables", len(d.sections), "sections")
        print("\n".join(out[:75]))

if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    inspect()
    if "--fetch" in sys.argv:
        for url,dest in [
            ("https://api.github.com/repos/Law-AI/summarization/git/trees/main?recursive=1", "research/inabs-tree.json"),
            ("https://api.github.com/repos/The-Atticus-Project/cuad/git/trees/main?recursive=1", "research/cuad-tree.json"),
            ("https://raw.githubusercontent.com/Law-AI/summarization/main/dataset/README.md", "research/inabs-readme.md"),
            ("https://huggingface.co/api/models/nlpaueb/legal-bert-small-uncased", "research/legalbert-model.json"),
            ("https://huggingface.co/api/models/google-t5/t5-small", "research/t5-model.json"),
        ]:
            try: fetch(url,dest)
            except Exception as exc: print(url,type(exc).__name__,str(exc))
